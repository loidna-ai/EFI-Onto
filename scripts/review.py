# -*- coding: utf-8 -*-
"""efi_tbox.ttl 의 판단 기준을 조사관 검토용 xlsx 로 내보낸다."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

G = json.load(open('build/graph.json', encoding='utf-8'))
O = json.load(open('build/onto.json', encoding='utf-8'))
KO = {n['id']: n['ko'] for n in G['nodes']}
k = lambda x: KO.get(x, x)
SCEN = {'PoorContactScenario': '접촉불량', 'CrushDamageScenario': '압착손상',
        'PartialDisconnectionScenario': '반단선', 'InsulationDegradationScenario': '절연열화',
        'TrackingScenario': '트래킹', 'ExternalFlameScenario': '외부화염',
        'ElectricalIgnitionScenario': '전기적 발화 공통'}
ROLE = {'Core': '핵심', 'Supporting': '보강', 'Refuting': '반증', 'DecisiveRefuting': '결정적 반증'}

FONT = 'Arial'
H_FILL = PatternFill('solid', fgColor='1F3A4D')
H_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
BODY = Font(name=FONT, size=10)
HI = PatternFill('solid', fgColor='FFF2CC')      # 우선 검토
IN = PatternFill('solid', fgColor='FFFFCC')      # 조사관 기입란
THIN = Side(style='thin', color='D0D0D0')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CTR = Alignment(horizontal='center', vertical='center')

wb = Workbook()

# ── 0. 안내 ────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = '검토 안내'
ws.sheet_view.showGridLines = False

DV = lambda: DataValidation(type='list', formula1='"O,X,△"', allow_blank=True)


def sheet(title, headers, data, widths, note, hi_idx=None, extra_inputs=()):
    w = wb.create_sheet(title)
    w.sheet_view.showGridLines = False
    w['A1'] = note
    w['A1'].font = Font(name=FONT, size=9, italic=True, color='666666')
    w.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers) + 2)
    w.row_dimensions[1].height = 26
    w['A1'].alignment = WRAP
    hdr = headers + list(extra_inputs) + ['판정', '의견·수정할 내용']
    for j, h in enumerate(hdr, 1):
        c = w.cell(3, j, h); c.font = H_FONT; c.fill = H_FILL; c.alignment = CTR; c.border = BOX
    dv = DV(); w.add_data_validation(dv)
    for i, row in enumerate(data, 4):
        for j, v in enumerate(row, 1):
            c = w.cell(i, j, v); c.font = BODY; c.border = BOX; c.alignment = WRAP
            if hi_idx is not None and hi_idx(row):
                c.fill = HI
        jd = len(headers) + 1
        ne = len(extra_inputs)
        for off in range(ne + 1):
            c = w.cell(i, jd + off); c.fill = IN; c.border = BOX; c.alignment = CTR
        w.cell(i, jd + ne + 1).fill = IN; w.cell(i, jd + ne + 1).border = BOX
        dv.add(w.cell(i, jd + ne))
    for j, wd in enumerate(widths + [8] * len(extra_inputs) + [7, 34], 1):
        w.column_dimensions[get_column_letter(j)].width = wd
    w.freeze_panes = 'A4'
    return w


# ── 1. 단서와 점수 ────────────────────────────────────────────────────────
rules = sorted(O['rules'], key=lambda r: (list(SCEN).index(r['sc'][0]) if r['sc'] and r['sc'][0] in SCEN else 9,
                                          -abs(r['d'][0] if r['d'] else 0)))
d1 = []
for r in rules:
    sc = SCEN.get(r['sc'][0], r['sc'][0]) if r['sc'] else ''
    ind = k(r['ind'][0]) if r['ind'] else ''
    role = ROLE.get(r['role'][0], '') if r['role'] else ''
    dv_ = r['d'][0] if r['d'] else 0
    st = '없음이 확인됨' if (r['st'] and r['st'][0] == 'ConfirmedAbsent') else '확인됨'
    sent = (f'{sc} 가설을 검토할 때, {ind}이(가) {st} 상태이면 '
            f'{"지지점수를 " + str(dv_) + "점 올린다" if dv_ > 0 else "지지점수를 " + str(-dv_) + "점 내린다"}.')
    d1.append([sc, ind, role, st, dv_, sent])
sheet('1 단서와 점수',
      ['가설', '확인 항목', '역할', '필요 상태', '가감점', '문장으로 읽으면'],
      d1, [12, 24, 12, 13, 8, 52],
      '가감점은 사람이 정한 값이 아니라 변별력에서 계산된 값입니다. 몇 가설이 같은 단서를 낼 수 있는지로 '
      '정해지며, 6개 가설 전부가 내는 단서는 0점이 됩니다. 그러니 점수 하나하나가 맞는지는 보지 마시고, '
      '같은 가설 안에서 항목들의 순서가 뒤집힌 곳이 있는지만 봐 주십시오. '
      '“이게 저것보다 중요한데 점수가 낮다” 싶은 곳을 X 로 표시해 주시면 됩니다.',
      hi_idx=lambda r: abs(r[4]) >= 15)

# ── 2. 형태 발현 ──────────────────────────────────────────────────────────
CH = O['chain']
shared = O['shared']
d2 = []
for a, b in sorted(CH['producesDamage'], key=lambda x: (k(x[0]), k(x[1]))):
    n = shared.get(b, 0)
    d2.append([k(a), k(b), n if n else '',
               f'{k(a)}이(가) 일어나면 {k(b)}이(가) 남을 수 있다.'])
sheet('2 형태 발현',
      ['발열 메커니즘', '남길 수 있는 손상 양상', '이 흔적을 남기는 가설 수', '문장으로 읽으면'],
      d2, [20, 24, 12, 52],
      '이 시트가 혼동 쌍 판정의 근거입니다. 여러 가설이 같은 흔적을 남기면 사진만으로 못 가른다고 판단합니다. '
      '노란 줄은 3개 이상의 가설이 공유하는 흔적으로, 변별력이 없다고 본 것입니다.',
      hi_idx=lambda r: isinstance(r[2], int) and r[2] >= 3)

# ── 3. 선행조건과 착화물 ─────────────────────────────────────────────────
d3 = []
for a, b in sorted(CH['enables'], key=lambda x: (k(x[0]), k(x[1]))):
    d3.append(['선행조건 → 발열', k(a), k(b), f'{k(a)} 상태에서는 {k(b)}이(가) 일어날 수 있다.'])
for a, b in sorted(CH['ignites'], key=lambda x: (k(x[0]), k(x[1]))):
    d3.append(['발열 → 착화물', k(a), k(b), f'{k(a)}은(는) {k(b)}에 불을 붙일 수 있다.'])
sheet('3 선행조건과 착화물', ['관계', '앞', '뒤', '문장으로 읽으면'], d3, [17, 22, 22, 52],
      '조건이 발열을 일으킬 수 있는지, 그 발열이 그 물질에 착화시킬 수 있는지를 봐 주십시오. '
      '특히 착화물 쪽은 실무 판단이 갈릴 수 있는 부분입니다.')

# ── 4. 현장사실과 증거물 ─────────────────────────────────────────────────
d4 = []
for a, b in sorted(CH['attests'], key=lambda x: (k(x[0]), k(x[1]))):
    d4.append(['현장사실 → 선행조건', k(a), k(b), f'{k(a)}이(가) 확인되면 {k(b)}이(가) 있었다고 본다.'])
for a, b in sorted(CH['exhibits'], key=lambda x: (k(x[0]), k(x[1]))):
    d4.append(['증거물 → 손상 양상', k(a), k(b), f'{k(a)}에서는 {k(b)}이(가) 관찰될 수 있다.'])
sheet('4 현장사실과 증거물', ['관계', '앞', '뒤', '문장으로 읽으면'], d4, [19, 22, 22, 52],
      '현장에서 확인한 사실이 어떤 조건의 존재를 뒷받침하는지, 각 증거물에서 어떤 흔적이 관찰될 수 있는지를 봐 주십시오.')

# ── 5. 용어 ───────────────────────────────────────────────────────────────
AXN = {'mech': '발열 메커니즘', 'evid': '물리 증거물', 'ante': '선행 조건', 'fuel': '착화물',
       'dmg': '손상 양상', 'scene': '현장 사실', 'scen': '가설·판정', 'evt': '사건',
       'obs': '계측', 'meta': '절차 기록', 'enum': '열거값', 'agent': '행위자', 'other': '기타'}
AX = {n['id']: n['ax'] for n in G['nodes']}
d5 = []
for n in sorted(G['nodes'], key=lambda x: (AXN.get(x['ax'], ''), x['ko'])):
    sup = O['c'].get(n['id'], {}).get('sup', [])
    d5.append([AXN.get(n['ax'], n['ax']), n['ko'], ' · '.join(k(s) for s in sup),
               O['c'].get(n['id'], {}).get('cm', '')])
sheet('5 용어', ['축', '이름', '상위 개념', '설명'], d5, [15, 24, 22, 48],
      '용어가 실무에서 쓰는 말과 맞는지, 상위-하위 관계가 뒤집힌 곳은 없는지 봐 주십시오. '
      '이름만 고치는 것은 결과에 영향을 주지 않으니 부담 없이 지적해 주셔도 됩니다.')


# ── 6. 반증 후보 ──────────────────────────────────────────────────────────
#  혼동 쌍마다 질문 하나. 후보 산출은 scripts/refute_audit.py 와 같은 함수다 —
#  거기서 사례 반례를 세고 여기서 원문·경험을 묻는다. 두 곳이 갈리면 안 된다.
#  이 시트는 원문 대조(audit·cause_audit)와 사례 반례 검사(make refute)를
#  거친 뒤에만 사람에게 간다. 묻는 것은 의견이 아니라 반례다.
import refute_audit
from rdflib import Graph
_g = Graph().parse('ontology/efi_tbox.ttl', format='turtle')
d6 = []
for sc, own, ind, n, morph in refute_audit.candidates(_g):
    d6.append([SCEN[sc], SCEN[own], k(ind), n, -21 if morph else -30,
               f'{k(ind)}은(는) 지금 {SCEN[own]}에서만 나온다고 되어 있습니다. '
               f'{SCEN[sc]} 화재에서 {k(ind)}이(가) 확인된 사건을 보신 적이 있습니까? '
               f'있으면 사건번호를 적어 주십시오.'])
d6.sort(key=lambda r: (r[0], r[4]))
sheet('6 반증 후보', ['가설', '경쟁 가설', '확인 항목', '공유 흔적 수', '가감점', '문장으로 읽으면'],
      d6, [12, 12, 22, 10, 8, 56],
      '사진으로 잘 안 갈리는 가설 쌍마다 질문을 하나씩 만들었습니다. 맞는지 틀린지가 아니라 '
      '반례를 묻습니다 — 그 요인의 화재에서 이 항목이 나온 사건을 보셨으면 사건번호를 적어 주십시오. '
      '반례 하나가 찬성 열 개보다 값집니다. 본 적이 없으면 비워 두셔도 됩니다. '
      '사건번호가 적힌 항목은 반증 후보에서 제외됩니다. 원문·사례 검증은 이미 거쳤습니다.')

# ── 7. 화재가 낼 수 있는 흔적인가 ────────────────────────────────────────
#  이 온톨로지는 '이 가설이 이 흔적을 낸다'는 방향으로만 지어졌다. 반대 방향을
#  빠뜨려 화재 자체가 만드는 흔적을 전기적 원인의 전용 단서로 오인한 사례가
#  원문 대조에서 다섯 건 나왔다. 남은 것을 조사관에게 묻는다.
import audit
_own = audit.owners(_g)
d7 = []
for pat in audit.unreviewed(_g):
    scs = sorted(SCEN.get(x, x) for x in _own.get(pat, ()))
    if not scs:
        continue
    d7.append([k(pat), ' · '.join(scs), len(scs),
               f'{k(pat)}은(는) 지금 {" · ".join(scs)}에서만 나타나는 흔적으로 되어 있습니다. '
               f'화재 자체의 열로도 이 흔적이 생길 수 있습니까?'])
d7.sort(key=lambda r: (-r[2], r[0]))
sheet('7 화재도 내는 흔적인가', ['손상 양상', '지금 이 가설들만 낸다고 되어 있음', '가설 수', '문장으로 읽으면'],
      d7, [24, 26, 8, 56],
      '화재 자체의 열로도 생기는 흔적이라면 그것은 전기적 원인의 전용 단서가 될 수 없습니다. '
      'O = 화재로도 생긴다, X = 화재로는 생기지 않는다. '
      'O 로 표시된 항목은 변별력 계산에서 빠지고 관련 가설의 점수가 내려갑니다. '
      '이미 같은 이유로 다섯 건을 고쳤습니다 — 탄화 도전로, 다발성 아크 비드, 아크흔, 변색, 광택 소실.')

wb.save('build/EFI-Onto_검토표.xlsx')
print('rows', len(d1), len(d2), len(d3), len(d4), len(d5), len(d6), len(d7))
